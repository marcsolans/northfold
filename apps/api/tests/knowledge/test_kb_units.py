"""Unit tests: hash embedder, chunking, parsing, SQL checks, column types, masking."""

from __future__ import annotations

import io
import math
from datetime import UTC, date, datetime
from decimal import Decimal
from email.message import EmailMessage
from typing import Any

import pytest
from kb_helpers import make_docx, make_pdf

from messer_common.knowledge import coltypes as ct
from messer_common.knowledge import policies as pol
from messer_common.knowledge.chunking import chunk_blocks, estimate_tokens
from messer_common.knowledge.embeddings import (
    HASH_MODEL,
    HashEmbedder,
    default_embedding_model,
    get_embedder,
    sigv4_headers,
    to_pgvector,
)
from messer_common.knowledge.errors import ParseError, PolicyError, SqlRejected
from messer_common.knowledge.parsing import Block, detect_kind, parse_sync
from messer_common.knowledge.sqlcheck import analyze, replace_spans, tokenize, validate_expression

OWN = "kb_0123456789abcdef0123456789abcdef"


def _cos(a: list[float], b: list[float]) -> float:
    return sum(x * y for x, y in zip(a, b, strict=True))


# ---------------------------------------------------------------------------------------------- embeddings


async def test_hash_embedder_is_deterministic_and_normalized() -> None:
    embedder = get_embedder(HASH_MODEL)
    assert embedder.dim == 384
    first = await embedder.embed(["Política de devoluciones: 30 días naturales"])
    second = await HashEmbedder().embed(["Política de devoluciones: 30 días naturales"])
    assert first == second
    assert math.isclose(math.sqrt(sum(x * x for x in first[0])), 1.0, rel_tol=1e-9)
    # Accents and case do not change the vector.
    assert await embedder.embed_query("CANCIÓN") == await embedder.embed_query("cancion")
    # Empty text still yields a unit vector (cosine distance stays defined).
    empty = await embedder.embed_query("")
    assert math.isclose(sum(x * x for x in empty), 1.0)


async def test_hash_embedder_similarity_reflects_overlap() -> None:
    e = HashEmbedder()
    query = await e.embed_query("plazo de devolución de productos")
    related = await e.embed_query("El plazo para la devolución de los productos es de 30 días")
    unrelated = await e.embed_query("Horario de la cafetería de la oficina central")
    assert _cos(query, related) > _cos(query, unrelated) + 0.2


async def test_default_model_respects_provider_and_residency(monkeypatch: pytest.MonkeyPatch) -> None:
    from messer_common.settings import get_settings

    settings = get_settings()
    assert await default_embedding_model() == HASH_MODEL  # tests run with MESSER_EMBEDDINGS_PROVIDER=hash
    monkeypatch.setattr(settings, "messer_embeddings_model", "voyage/voyage-3.5")
    assert await default_embedding_model() == "voyage/voyage-3.5"
    assert await default_embedding_model(data_residency="eu_strict") == HASH_MODEL


def test_pgvector_literal_and_sigv4_headers() -> None:
    assert to_pgvector([0.5, -1.0, 1e-9]) == "[0.5,-1,1e-09]"
    moment = datetime(2026, 1, 2, 3, 4, 5, tzinfo=UTC)
    a = sigv4_headers(
        method="POST",
        url="https://bedrock-runtime.eu-central-1.amazonaws.com/model/amazon.titan-embed-text-v2%3A0/invoke",
        region="eu-central-1",
        service="bedrock",
        body=b"{}",
        access_key="AKIDEXAMPLE",
        secret_key="secret",
        now=moment,
    )
    b = sigv4_headers(
        method="POST",
        url="https://bedrock-runtime.eu-central-1.amazonaws.com/model/amazon.titan-embed-text-v2%3A0/invoke",
        region="eu-central-1",
        service="bedrock",
        body=b"{}",
        access_key="AKIDEXAMPLE",
        secret_key="secret",
        now=moment,
    )
    assert a == b
    assert a["X-Amz-Date"] == "20260102T030405Z"
    assert a["Authorization"].startswith(
        "AWS4-HMAC-SHA256 Credential=AKIDEXAMPLE/20260102/eu-central-1/bedrock/aws4_request, "
        "SignedHeaders=content-type;host;x-amz-date, Signature="
    )


# ---------------------------------------------------------------------------------------------- chunking


def test_estimate_tokens_heuristic() -> None:
    assert estimate_tokens("") == 0
    assert estimate_tokens("hola mundo") == 2
    assert estimate_tokens("internacionalización") == 4
    assert estimate_tokens("¿Qué tal?") == 4


def test_chunking_keeps_structure_pages_and_overlap() -> None:
    paragraph = " ".join(f"Frase número {i} sobre la política de garantías del producto." for i in range(40))
    blocks = [
        Block("heading", "Manual", level=1, page=1),
        Block("heading", "Garantía", level=2, page=1),
        Block("text", paragraph, page=1),
        Block("text", paragraph, page=2),
        Block("heading", "Devoluciones", level=2, page=3),
        Block("text", "Las devoluciones se aceptan durante 30 días. " * 10, page=3),
    ]
    chunks = chunk_blocks(blocks, max_tokens=300, overlap_tokens=60)
    assert len(chunks) >= 4
    assert all(c.token_count <= 300 for c in chunks)
    assert [c.ordinal for c in chunks] == list(range(len(chunks)))
    garantia = [c for c in chunks if c.heading == "Manual > Garantía"]
    devoluciones = [c for c in chunks if c.heading == "Manual > Devoluciones"]
    assert len(garantia) >= 3 and len(devoluciones) == 1
    assert devoluciones[0].page_from == 3 and devoluciones[0].page_to == 3
    assert any(c.page_from == 1 and c.page_to == 2 for c in garantia) or {c.page_from for c in garantia} == {1, 2}
    # Consecutive chunks of the same section share trailing context.
    second = garantia[1]
    assert second.metadata["overlap_tokens"] > 0
    tail = garantia[0].content.split(". ")[-2]
    assert tail.split()[-1] in second.content
    # Sections never inherit overlap across a heading.
    assert devoluciones[0].metadata["overlap_tokens"] == 0
    assert "Frase número" not in devoluciones[0].content


def test_chunking_merges_tiny_sections() -> None:
    blocks = [Block("heading", f"Paso {i}", level=2) for i in range(1)]
    blocks += [Block("text", "Abrir la caja."), Block("heading", "Paso 2", level=2), Block("text", "Enchufar.")]
    chunks = chunk_blocks(blocks, max_tokens=800, overlap_tokens=120)
    assert len(chunks) == 1
    assert "## Paso 2" in chunks[0].content and chunks[0].heading == "Paso 0"


# ---------------------------------------------------------------------------------------------- parsing


def test_pdf_parsing_detects_headings_and_pages() -> None:
    data = make_pdf(
        [
            [("Informe anual", 24), ("Resumen", 16), ("Las ventas crecieron un 12 % en Europa.", 11)],
            [("Riesgos", 16), ("El principal riesgo es la divisa.", 11)],
        ]
    )
    parsed = parse_sync(data, "informe.pdf", "application/pdf")
    assert parsed.kind == "pdf" and parsed.pages == 2
    headings = [(b.text, b.level, b.page) for b in parsed.blocks if b.kind == "heading"]
    assert headings == [("Informe anual", 1, 1), ("Resumen", 2, 1), ("Riesgos", 2, 2)]
    texts = [b for b in parsed.blocks if b.kind == "text"]
    assert texts[-1].text == "El principal riesgo es la divisa." and texts[-1].page == 2


def test_docx_parsing_headings_lists_tables() -> None:
    import docx

    data = make_docx([("h1", "Contrato"), ("p", "Las partes acuerdan."), ("h2", "Pago"), ("li", "Mensual")])
    document = docx.Document(io.BytesIO(data))
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text, table.cell(0, 1).text = "Concepto", "Importe"
    table.cell(1, 0).text, table.cell(1, 1).text = "Licencia", "100"
    buf = io.BytesIO()
    document.save(buf)
    parsed = parse_sync(buf.getvalue(), "contrato.docx")
    kinds = [(b.kind, b.text, b.level) for b in parsed.blocks]
    assert ("heading", "Contrato", 1) in kinds and ("heading", "Pago", 2) in kinds
    assert ("text", "- Mensual", 0) in kinds
    assert ("table", "Concepto | Importe\nLicencia | 100", 0) in kinds


def test_pptx_parsing_slides_as_pages() -> None:
    from pptx import Presentation

    prs = Presentation()
    slide: Any = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = "Objetivos 2026"
    slide.placeholders[1].text_frame.text = "Crecer en Alemania"
    buf = io.BytesIO()
    prs.save(buf)
    parsed = parse_sync(buf.getvalue(), "plan.pptx")
    assert parsed.pages == 1
    assert parsed.blocks[0].kind == "heading" and parsed.blocks[0].text == "Objetivos 2026"
    assert any(b.text == "Crecer en Alemania" and b.page == 1 for b in parsed.blocks)


def test_csv_and_xlsx_become_tables() -> None:
    csv_data = "Nombre;Importe;Fecha\nAcme;1.234,50;31/12/2025\nBeta;10,00;01/01/2026\n".encode("cp1252")
    parsed = parse_sync(csv_data, "facturas.csv")
    assert parsed.metadata["delimiter"] == ";"
    table = parsed.tables[0]
    assert table.name == "facturas" and table.columns == ["Nombre", "Importe", "Fecha"]
    assert ct.infer_columns(table.columns, table.rows) == ["text", "numeric", "date"]

    import openpyxl

    wb = openpyxl.Workbook()
    ws: Any = wb.active
    ws.title = "Clientes"
    ws.append(["id", "alta", "activo"])
    ws.append([1, datetime(2025, 3, 1), True])
    ws.append([2, datetime(2025, 4, 1), False])
    buf = io.BytesIO()
    wb.save(buf)
    parsed = parse_sync(buf.getvalue(), "crm.xlsx")
    table = parsed.tables[0]
    assert table.sheet == "Clientes"
    assert ct.infer_columns(table.columns, table.rows) == ["integer", "date", "boolean"]


def test_html_markdown_and_email() -> None:
    html = (
        b"<html><head><title>Ayuda</title><style>p{}</style></head><body><h1>FAQ</h1><p>Hola <b>mundo</b></p>"
        b"<script>alert(1)</script><ul><li>Uno</li></ul><table><tr><td>a</td><td>b</td></tr></table></body></html>"
    )
    parsed = parse_sync(html, "faq.html")
    assert parsed.metadata["title"] == "Ayuda"
    assert [(b.kind, b.text) for b in parsed.blocks] == [
        ("heading", "FAQ"),
        ("text", "Hola mundo"),
        ("text", "- Uno"),
        ("table", "a | b"),
    ]
    md = b"---\ntitle: Guia\n---\n# Inicio\n\nTexto inicial.\n\n## Detalle\n\n```sql\nselect 1;\n```\n"
    parsed = parse_sync(md, "guia.md")
    assert parsed.metadata == {"title": "Guia"}
    assert [(b.kind, b.text, b.level) for b in parsed.blocks] == [
        ("heading", "Inicio", 1),
        ("text", "Texto inicial.", 0),
        ("heading", "Detalle", 2),
        ("code", "select 1;", 0),
    ]
    msg = EmailMessage()
    msg["Subject"] = "Pedido 123"
    msg["From"] = "ana@acme.es"
    msg["To"] = "ventas@messer.local"
    msg.set_content("Necesito el presupuesto.\n\nGracias.")
    msg.add_attachment(b"Condiciones: pago a 30 dias.", maintype="text", subtype="plain", filename="condiciones.txt")
    parsed = parse_sync(msg.as_bytes(), "pedido.eml")
    texts = [b.text for b in parsed.blocks]
    assert texts[0] == "Pedido 123" and "De: ana@acme.es" in texts[1]
    assert "Adjunto: condiciones.txt" in texts and "Condiciones: pago a 30 dias." in texts
    assert parsed.metadata["attachments"] == ["condiciones.txt"]


def test_unsupported_formats_are_rejected() -> None:
    with pytest.raises(ParseError) as err:
        detect_kind("viejo.doc")
    assert "docx" in err.value.message
    with pytest.raises(ParseError):
        detect_kind("foto.png", "image/png", b"\x89PNG")
    assert detect_kind("sin-extension", None, b"%PDF-1.7") == "pdf"


# ---------------------------------------------------------------------------------------------- SQL checks


@pytest.mark.parametrize(
    "sql",
    [
        "SELECT * FROM clientes",
        "select 1;",
        "WITH t AS (SELECT 1 AS a) SELECT a FROM t",
        "EXPLAIN (ANALYZE, FORMAT JSON) SELECT * FROM clientes",
        "VALUES (1), (2)",
        "TABLE clientes",
        "(SELECT 1) UNION ALL (SELECT 2)",
        "SELECT 'a;b' AS x, $$; DROP TABLE x$$ AS y -- ; comment\n",
        "SELECT * FROM clientes ORDER BY 1 FETCH FIRST 10 ROWS ONLY",
        "SELECT substring(nombre FROM 1 FOR 3) FROM clientes",
        "SELECT pg_typeof(1), information_schema._pg_truetypid(a, t) FROM information_schema.columns a, x t",
        f'SELECT * FROM "{OWN}".clientes',
        "SELECT chunk_id FROM kb_vector_search($1::uuid, embed('garantía'), 5)",
    ],
)
def test_sqlcheck_accepts_read_only(sql: str) -> None:
    analyze(sql, own_schema=OWN)


@pytest.mark.parametrize(
    ("sql", "code"),
    [
        ("DELETE FROM clientes", "sql_not_read_only"),
        ("SELECT 1; SELECT 2", "sql_multiple_statements"),
        ("SELECT 1; DROP TABLE clientes", "sql_multiple_statements"),
        ("SELECT * INTO copia FROM clientes", "sql_not_read_only"),
        ("WITH d AS (DELETE FROM clientes RETURNING *) SELECT * FROM d", "sql_not_read_only"),
        ("SELECT * FROM clientes FOR UPDATE", "sql_not_read_only"),
        ("SELECT set_config('role', 'messer_readonly', true)", "sql_forbidden_function"),
        ("SELECT \"set_config\"('app.kb_scope', 'x', true)", "sql_forbidden_function"),
        ("SELECT pg_catalog.set_config('a', 'b', true)", "sql_forbidden_function"),
        ("SELECT query_to_xml('select 1', true, false, '')", "sql_forbidden_function"),
        ("SELECT pg_read_file('/etc/passwd')", "sql_forbidden_function"),
        ("SELECT pg_sleep(10)", "sql_forbidden_function"),
        ("SELECT lo_import('/etc/passwd')", "sql_forbidden_function"),
        ("SELECT * FROM pg_catalog.pg_class", "sql_forbidden"),
        ("SELECT * FROM pg_roles", "sql_forbidden"),
        ("SELECT * FROM kb_ffffffffffffffffffffffffffffffff.clientes", "sql_scope"),
        ("SELECT * FROM ONLY clientes", "sql_forbidden"),
        ("UPDATE clientes SET nombre = 'x'", "sql_not_read_only"),
        ("CREATE TABLE x (a int)", "sql_not_read_only"),
        ("COPY clientes TO STDOUT", "sql_not_read_only"),
        ("SET ROLE messer", "sql_not_read_only"),
        ("SELECT U&\"\\0073et_config\"('a','b',true)", "sql_forbidden"),
        ("SELECT 'sin cerrar", "sql_unterminated"),
        ("", "sql_empty"),
        ("SELECT embed(nombre) FROM clientes", "sql_embed"),
    ],
)
def test_sqlcheck_rejects(sql: str, code: str) -> None:
    with pytest.raises(SqlRejected) as err:
        analyze(sql, own_schema=OWN)
    assert err.value.code == code


def test_sqlcheck_collects_params_names_and_embed_calls() -> None:
    sql = "SELECT c.nombre FROM clientes c JOIN kb_text_search($1::uuid, $2, 5) t ON true WHERE embed($2) IS NOT NULL"
    a = analyze(sql, own_schema=OWN)
    assert a.max_param == 2 and a.params_used == {1, 2}
    assert "clientes" in a.names and "kb_text_search" in a.functions
    assert len(a.embed_calls) == 1 and a.embed_calls[0].param == 2
    rewritten = replace_spans(sql, [(a.embed_calls[0].start, a.embed_calls[0].end, "$3::vector")])
    assert rewritten.endswith("WHERE $3::vector IS NOT NULL")
    assert [t.value for t in tokenize("E'it''s\\n' \"Col\"\"x\"")] == ["it's\n", 'Col"x']


def test_policy_filter_validation() -> None:
    validate_expression("region = 'norte' AND importe < 1000")
    validate_expression("metadata->>'confidencial' IS DISTINCT FROM 'true'")
    for bad in (
        "1=1; DROP TABLE x",
        "id IN (SELECT id FROM otra)",
        "set_config('a','b',true) IS NOT NULL",
        "t.region = 'x'",
    ):
        with pytest.raises(SqlRejected):
            validate_expression(bad)


# ---------------------------------------------------------------------------------------------- types & masks


def test_type_inference_and_coercion() -> None:
    assert ct.infer_type(["1", "2", "-3"]) == "integer"
    assert ct.infer_type(["00123", "456"]) == "text"  # leading zeros: identifiers, postal codes
    assert ct.infer_type(["1.234,56", "7,5", ""]) == "numeric"
    assert ct.infer_type(["1,234.56", "10.5"]) == "numeric"
    assert ct.infer_type(["sí", "no", "Si"]) == "boolean"
    assert ct.infer_type(["2025-01-31", "31/12/2024"]) == "date"
    assert ct.infer_type(["2025-01-31 10:00", "2025-02-01T08:30:00Z"]) == "timestamp"
    assert ct.infer_type(['{"a": 1}', "[1, 2]"]) == "json"
    assert ct.infer_type(["B12345678", "A87654321"]) == "text"
    assert ct.coerce("1.234,56", "numeric") == Decimal("1234.56")
    assert ct.coerce("31/12/2024", "date") == date(2024, 12, 31)
    assert ct.coerce("sí", "boolean") is True
    assert ct.coerce("", "integer") is None
    with pytest.raises(ct.CoercionError):
        ct.coerce("abc", "integer")
    assert ct.unique_identifiers(["Nombre Cliente", "nombre-cliente", "Año", "select", "1º trimestre", "_row_id"]) == [
        "nombre_cliente",
        "nombre_cliente_2",
        "ano",
        "select_",
        "c_1o_trimestre",
        "row_id",
    ]


def test_masks_and_policy_combination() -> None:
    assert pol.mask_value("ana@acme.es", "redact") == "***"
    assert pol.mask_value("ana@acme.es", "hash") == pol.mask_value("ana@acme.es", "hash")
    assert len(pol.mask_value("ana@acme.es", "hash")) == 16
    assert pol.mask_value("B12345678", "truncate:3") == "B12"
    assert pol.mask_value(None, "redact") is None
    with pytest.raises(PolicyError):
        pol.parse_mask("encrypt")
    assert pol.stricter("truncate:4", "hash") == "hash"
    assert pol.stricter("truncate:4", "truncate:2") == "truncate:2"
    assert pol.policy_role(None) == "agent" and pol.policy_role("viewer") == "viewer"
