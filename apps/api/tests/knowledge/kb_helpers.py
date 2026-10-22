"""Helpers for the knowledge tests: generated documents, multipart bodies, running the ingestion job."""

from __future__ import annotations

import io
import uuid
from typing import Any

from sqlalchemy import text

from messer_common.db import system_session


def make_pdf(pages: list[list[tuple[str, float]]]) -> bytes:
    """PDF with one list of (text, font size) lines per page."""
    import fitz

    doc: Any = fitz.open()
    for lines in pages:
        page = doc.new_page()
        y = 72.0
        for line, size in lines:
            page.insert_text((72, y), line, fontsize=size)
            y += size * 1.8
    data = doc.tobytes()
    doc.close()
    return bytes(data)


def make_docx(blocks: list[tuple[str, str]]) -> bytes:
    """DOCX from (style, text) blocks; style "h1"/"h2" or "p" or "li"."""
    import docx

    document: Any = docx.Document()
    for style, content in blocks:
        if style.startswith("h"):
            document.add_heading(content, level=int(style[1]))
        elif style == "li":
            document.add_paragraph(content, style="List Bullet")
        else:
            document.add_paragraph(content)
    buf = io.BytesIO()
    document.save(buf)
    return buf.getvalue()


def multipart(*files: tuple[str, bytes, str]) -> list[tuple[str, tuple[str, bytes, str]]]:
    return [("files", (name, data, mime)) for name, data, mime in files]


async def ingest_all(kb_id: uuid.UUID) -> list[str]:
    """Run the worker job for every queued file of a knowledge base (no worker process in tests)."""
    from messer_worker.jobs.knowledge import ingest_file

    async with system_session() as s:
        rows = await s.execute(
            text("SELECT id FROM public.kb_files WHERE kb_id = :kb AND status = 'queued' ORDER BY created_at"),
            {"kb": kb_id},
        )
        ids = [str(r[0]) for r in rows]
    return [await ingest_file({"job_try": 1}, file_id) for file_id in ids]
