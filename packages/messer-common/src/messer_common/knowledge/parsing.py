"""Document parsing into structured blocks (headings, paragraphs, tables) and tabular data.

Docling (MIT) is used for PDF, DOCX, PPTX and HTML when it is installed and ``MESSER_KB_DOCLING_ENABLED``
is on. It is an optional adapter: Docling pulls PyTorch and model weights, which conflicts with the
engine-pinned shared environment, so the default deployment uses the built-in parsers below. Any Docling
failure falls back to them.

Built-in parsers: PDF (pymupdf: pages, headings detected by font size), DOCX (python-docx: heading styles,
lists, tables), PPTX (slide titles, text frames, tables, notes), XLSX and CSV (→ tables), HTML (tags
stripped, h1–h6 headings), Markdown/TXT and e-mail (.eml, with parsable attachments).
"""

from __future__ import annotations

import asyncio
import csv
import email
import email.policy
import io
import json
import re
import tempfile
from collections import Counter
from dataclasses import dataclass, field
from email.message import EmailMessage
from html.parser import HTMLParser
from pathlib import Path
from typing import Any, Literal

from ..logging import get_logger
from ..settings import get_settings
from .errors import ParseError

log = get_logger("messer.knowledge.parsing")

BlockKind = Literal["heading", "text", "table", "code"]
MAX_TABLE_ROWS = 1_000_000


@dataclass(slots=True)
class Block:
    kind: BlockKind
    text: str
    level: int = 0  # heading level (1 = top)
    page: int | None = None


@dataclass(slots=True)
class ParsedTable:
    name: str
    columns: list[str]
    rows: list[list[Any]]
    sheet: str | None = None


@dataclass(slots=True)
class ParsedDocument:
    kind: str
    parser: str
    blocks: list[Block] = field(default_factory=list)
    tables: list[ParsedTable] = field(default_factory=list)
    pages: int | None = None
    metadata: dict[str, Any] = field(default_factory=dict)

    @property
    def has_text(self) -> bool:
        return any(b.text.strip() for b in self.blocks)


EXTENSIONS: dict[str, str] = {
    ".pdf": "pdf",
    ".docx": "docx",
    ".pptx": "pptx",
    ".xlsx": "xlsx",
    ".xlsm": "xlsx",
    ".csv": "csv",
    ".tsv": "csv",
    ".html": "html",
    ".htm": "html",
    ".md": "markdown",
    ".markdown": "markdown",
    ".txt": "text",
    ".text": "text",
    ".log": "text",
    ".json": "text",
    ".xml": "text",
    ".eml": "eml",
}
MIME_KINDS: dict[str, str] = {
    "application/pdf": "pdf",
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document": "docx",
    "application/vnd.openxmlformats-officedocument.presentationml.presentation": "pptx",
    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet": "xlsx",
    "application/vnd.ms-excel.sheet.macroenabled.12": "xlsx",
    "text/csv": "csv",
    "text/tab-separated-values": "csv",
    "text/html": "html",
    "text/markdown": "markdown",
    "text/x-markdown": "markdown",
    "text/plain": "text",
    "application/json": "text",
    "message/rfc822": "eml",
}
UNSUPPORTED_HINTS = {
    ".doc": "Guarda el documento como .docx",
    ".xls": "Guarda la hoja de cálculo como .xlsx o .csv",
    ".ppt": "Guarda la presentación como .pptx",
}
TABULAR_KINDS = {"csv", "xlsx"}

MIME_BY_KIND = {
    "pdf": "application/pdf",
    "docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    "pptx": "application/vnd.openxmlformats-officedocument.presentationml.presentation",
    "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    "csv": "text/csv",
    "html": "text/html",
    "markdown": "text/markdown",
    "text": "text/plain",
    "eml": "message/rfc822",
}


def detect_kind(filename: str, mime: str | None = None, data: bytes | None = None) -> str:
    """Document kind from extension, then MIME type, then magic bytes."""
    suffix = Path(filename).suffix.lower()
    if suffix in EXTENSIONS:
        kind = EXTENSIONS[suffix]
        if suffix == ".tsv":
            return "csv"
        return kind
    if suffix in UNSUPPORTED_HINTS:
        raise ParseError(
            f"Formato {suffix} no soportado. {UNSUPPORTED_HINTS[suffix]}.",
            code="unsupported_format",
            details={"extension": suffix},
        )
    base_mime = (mime or "").split(";")[0].strip().lower()
    if base_mime in MIME_KINDS:
        return MIME_KINDS[base_mime]
    if data is not None:
        if data.startswith(b"%PDF"):
            return "pdf"
        if data[:5].lower().startswith(b"<!doc") or data[:6].lower().startswith(b"<html"):
            return "html"
    raise ParseError(
        "Formato de archivo no soportado",
        code="unsupported_format",
        details={"filename": filename, "mime": mime},
    )


def guess_mime(filename: str, mime: str | None = None) -> str:
    base = (mime or "").split(";")[0].strip().lower()
    if base and base != "application/octet-stream":
        return base
    try:
        return MIME_BY_KIND[detect_kind(filename, None)]
    except ParseError:
        return "application/octet-stream"


def decode_text(data: bytes) -> str:
    """UTF-8 (with or without BOM), then UTF-16 with BOM, then Windows-1252."""
    if data.startswith((b"\xff\xfe", b"\xfe\xff")):
        return data.decode("utf-16")
    try:
        return data.decode("utf-8-sig")
    except UnicodeDecodeError:
        return data.decode("cp1252", errors="replace")


def _clean(text: str) -> str:
    text = text.replace("\x00", "").replace("\r\n", "\n").replace("\r", "\n")
    text = re.sub(r"[ \t ]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


# ---------------------------------------------------------------------------------------------- PDF


def _parse_pdf(data: bytes) -> ParsedDocument:
    import fitz  # pymupdf

    try:
        doc: Any = fitz.open(stream=data, filetype="pdf")
    except Exception as exc:
        raise ParseError("El PDF está dañado o no se puede abrir", details={"error": str(exc)[:200]}) from exc
    if doc.needs_pass:
        raise ParseError("El PDF está protegido con contraseña", code="pdf_encrypted")
    lines: list[tuple[str, float, bool, int]] = []  # text, size, bold, page
    sizes: Counter[float] = Counter()
    for page_no, page in enumerate(doc, start=1):
        info: Any = page.get_text("dict", sort=True)
        for block in info.get("blocks", []):
            if block.get("type") != 0:
                continue
            parts: list[str] = []
            max_size = 0.0
            bold_chars = 0
            chars = 0
            for line in block.get("lines", []):
                line_text = "".join(span.get("text", "") for span in line.get("spans", []))
                for span in line.get("spans", []):
                    text = span.get("text", "")
                    n = len(text.strip())
                    if n == 0:
                        continue
                    size = round(float(span.get("size", 0.0)) * 2) / 2
                    sizes[size] += n
                    chars += n
                    max_size = max(max_size, size)
                    if int(span.get("flags", 0)) & 16:
                        bold_chars += n
                if line_text.strip():
                    parts.append(line_text.strip())
            text = _clean(" ".join(parts))
            if text:
                lines.append((text, max_size, chars > 0 and bold_chars / chars > 0.6, page_no))
    page_count = int(doc.page_count)
    metadata: dict[str, Any] = {
        k: v for k, v in (doc.metadata or {}).items() if isinstance(v, str) and v.strip() and k in ("title", "author")
    }
    doc.close()

    result = ParsedDocument(kind="pdf", parser="pymupdf", pages=page_count, metadata=metadata)
    if not lines:
        return result
    body = sizes.most_common(1)[0][0] if sizes else 10.0
    heading_sizes = sorted({s for _, s, _, _ in lines if s >= body * 1.15}, reverse=True)
    level_of = {s: min(i + 1, 4) for i, s in enumerate(heading_sizes)}
    for text, size, bold, page in lines:
        short = len(text) <= 180 and not text.endswith((".", ":", ";", ","))
        if size in level_of and short:
            result.blocks.append(Block("heading", text, level=level_of[size], page=page))
        elif bold and short and len(text) <= 100 and size >= body:
            result.blocks.append(Block("heading", text, level=len(heading_sizes) + 1, page=page))
        else:
            result.blocks.append(Block("text", text, page=page))
    return result


# ---------------------------------------------------------------------------------------------- DOCX

_HEADING_STYLE = re.compile(r"(heading|t[ií]tulo|encabezado|überschrift|titre)\s*(\d)", re.IGNORECASE)


def _docx_table_text(table: Any) -> str:
    rows: list[str] = []
    for row in table.rows:
        cells: list[str] = []
        for cell in row.cells:
            value = _clean(cell.text)
            if not cells or cells[-1] != value:  # merged cells repeat their text
                cells.append(value)
        if any(cells):
            rows.append(" | ".join(cells))
    return "\n".join(rows)


def _parse_docx(data: bytes) -> ParsedDocument:
    import docx
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    try:
        document: Any = docx.Document(io.BytesIO(data))
    except Exception as exc:
        raise ParseError(
            "El documento Word está dañado o no se puede abrir", details={"error": str(exc)[:200]}
        ) from exc
    result = ParsedDocument(kind="docx", parser="python-docx")
    props = document.core_properties
    if props.title:
        result.metadata["title"] = props.title
    if props.author:
        result.metadata["author"] = props.author
    body = document.element.body
    for child in body.iterchildren():
        tag = child.tag.rsplit("}", 1)[-1]
        if tag == "p":
            para = Paragraph(child, document)
            text = _clean(para.text)
            if not text:
                continue
            style = (para.style.name if para.style is not None else "") or ""
            match = _HEADING_STYLE.search(style)
            if style.lower() in ("title", "título", "titulo"):
                result.blocks.append(Block("heading", text, level=1))
            elif match:
                result.blocks.append(Block("heading", text, level=min(int(match.group(2)), 6)))
            elif "list" in style.lower() or "lista" in style.lower():
                result.blocks.append(Block("text", f"- {text}"))
            else:
                result.blocks.append(Block("text", text))
        elif tag == "tbl":
            text = _docx_table_text(Table(child, document))
            if text:
                result.blocks.append(Block("table", text))
    return result


# ---------------------------------------------------------------------------------------------- PPTX


def _parse_pptx(data: bytes) -> ParsedDocument:
    from pptx import Presentation

    try:
        prs: Any = Presentation(io.BytesIO(data))
    except Exception as exc:
        raise ParseError("La presentación está dañada o no se puede abrir", details={"error": str(exc)[:200]}) from exc
    result = ParsedDocument(kind="pptx", parser="python-pptx")
    slide_count = 0
    for index, slide in enumerate(prs.slides, start=1):
        slide_count = index
        title_shape = slide.shapes.title
        title = _clean(title_shape.text_frame.text) if title_shape is not None and title_shape.has_text_frame else ""
        result.blocks.append(Block("heading", title or f"Diapositiva {index}", level=1, page=index))
        for shape in slide.shapes:
            if title_shape is not None and shape.shape_id == title_shape.shape_id:
                continue
            if getattr(shape, "has_table", False) and shape.has_table:
                rows = [" | ".join(_clean(c.text) for c in row.cells) for row in shape.table.rows]
                text = "\n".join(r for r in rows if r.strip(" |"))
                if text:
                    result.blocks.append(Block("table", text, page=index))
            elif getattr(shape, "has_text_frame", False) and shape.has_text_frame:
                for paragraph in shape.text_frame.paragraphs:
                    text = _clean("".join(run.text for run in paragraph.runs))
                    if text:
                        prefix = "- " if paragraph.level > 0 else ""
                        result.blocks.append(Block("text", prefix + text, page=index))
        if slide.has_notes_slide:
            notes = _clean(slide.notes_slide.notes_text_frame.text if slide.notes_slide.notes_text_frame else "")
            if notes:
                result.blocks.append(Block("text", f"Notas: {notes}", page=index))
    result.pages = slide_count
    return result


# ---------------------------------------------------------------------------------------------- tables


def _unique_columns(raw: list[Any]) -> list[str]:
    names: list[str] = []
    seen: Counter[str] = Counter()
    for i, value in enumerate(raw, start=1):
        name = str(value).strip() if value is not None and str(value).strip() else f"column_{i}"
        seen[name] += 1
        names.append(name if seen[name] == 1 else f"{name}_{seen[name]}")
    return names


def _rows_to_table(name: str, raw_rows: list[list[Any]], sheet: str | None = None) -> ParsedTable | None:
    def empty(v: Any) -> bool:
        return v is None or (isinstance(v, str) and not v.strip())

    rows = [r for r in raw_rows if any(not empty(v) for v in r)]
    if not rows:
        return None
    width = max(len(r) for r in rows)
    # Drop trailing columns that are empty everywhere.
    while width > 0 and all(len(r) < width or empty(r[width - 1]) for r in rows):
        width -= 1
    if width == 0:
        return None
    padded = [list(r[:width]) + [None] * (width - len(r[:width])) for r in rows]
    header, data = padded[0], padded[1:]
    if len(data) > MAX_TABLE_ROWS:
        raise ParseError(
            f"La tabla supera el máximo de {MAX_TABLE_ROWS:,} filas".replace(",", "."), code="table_too_large"
        )
    return ParsedTable(name=name, columns=_unique_columns(header), rows=data, sheet=sheet)


def _sniff_delimiter(sample: str) -> str:
    try:
        return csv.Sniffer().sniff(sample, delimiters=",;\t|").delimiter
    except csv.Error:
        counts = {d: sample.count(d) for d in (";", ",", "\t", "|")}
        return max(counts, key=lambda d: counts[d]) if any(counts.values()) else ","


def _parse_csv(data: bytes, filename: str) -> ParsedDocument:
    text = decode_text(data)
    delimiter = "\t" if filename.lower().endswith(".tsv") else _sniff_delimiter(text[:65536])
    reader = csv.reader(io.StringIO(text), delimiter=delimiter)
    try:
        raw = [row for row in reader]
    except csv.Error as exc:
        raise ParseError("El CSV no tiene un formato válido", details={"error": str(exc)}) from exc
    result = ParsedDocument(kind="csv", parser="csv", metadata={"delimiter": delimiter})
    table = _rows_to_table(Path(filename).stem, [list(r) for r in raw])
    if table is not None:
        result.tables.append(table)
    return result


def _parse_xlsx(data: bytes, filename: str) -> ParsedDocument:
    import openpyxl

    try:
        workbook: Any = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as exc:
        raise ParseError("El Excel está dañado o no se puede abrir", details={"error": str(exc)[:200]}) from exc
    result = ParsedDocument(kind="xlsx", parser="openpyxl", metadata={"sheets": list(workbook.sheetnames)})
    stem = Path(filename).stem
    sheets = list(workbook.worksheets)
    for ws in sheets:
        rows: list[list[Any]] = []
        for row in ws.iter_rows(values_only=True):
            rows.append(list(row))
            if len(rows) > MAX_TABLE_ROWS + 1:
                raise ParseError("La hoja supera el máximo de filas admitido", code="table_too_large")
        name = stem if len(sheets) == 1 else f"{stem}_{ws.title}"
        table = _rows_to_table(name, rows, sheet=ws.title)
        if table is not None:
            result.tables.append(table)
    workbook.close()
    return result


# ---------------------------------------------------------------------------------------------- HTML


class _HtmlText(HTMLParser):
    BLOCK_TAGS = frozenset(
        "p div section article header footer main aside nav li ul ol tr table blockquote dd dt figcaption br hr form".split()
    )
    SKIP_TAGS = frozenset({"script", "style", "noscript", "template", "svg"})

    def __init__(self) -> None:
        super().__init__(convert_charrefs=True)
        self.blocks: list[Block] = []
        self.title = ""
        self._buf: list[str] = []
        self._skip = 0
        self._heading: int | None = None
        self._pre = 0
        self._in_title = False
        self._cells: list[str] | None = None
        self._in_head = False

    def _flush(self) -> None:
        text = _clean("".join(self._buf))
        self._buf = []
        if not text:
            return
        if self._heading is not None:
            self.blocks.append(Block("heading", text, level=self._heading))
        elif self._pre:
            self.blocks.append(Block("code", text))
        else:
            self.blocks.append(Block("text", text))

    def handle_starttag(self, tag: str, attrs: list[tuple[str, str | None]]) -> None:
        if tag == "head":
            self._in_head = True
            return
        if tag == "body":
            self._in_head = False
            return
        if tag in self.SKIP_TAGS:
            self._skip += 1
            return
        if tag == "title":
            self._in_title = True
            return
        if self._skip or self._in_head:
            return
        if re.fullmatch(r"h[1-6]", tag):
            self._flush()
            self._heading = int(tag[1])
        elif tag == "pre":
            self._flush()
            self._pre += 1
        elif tag == "tr":
            self._flush()
            self._cells = []
        elif tag in ("td", "th"):
            self._buf = []
        elif tag in self.BLOCK_TAGS:
            self._flush()
            if tag == "li":
                self._buf.append("- ")

    def handle_endtag(self, tag: str) -> None:
        if tag == "head":
            self._in_head = False
            return
        if tag in self.SKIP_TAGS:
            self._skip = max(0, self._skip - 1)
            return
        if tag == "title":
            self._in_title = False
            return
        if self._skip or self._in_head:
            return
        if re.fullmatch(r"h[1-6]", tag):
            self._flush()
            self._heading = None
        elif tag == "pre":
            self._flush()
            self._pre = max(0, self._pre - 1)
        elif tag in ("td", "th") and self._cells is not None:
            self._cells.append(_clean("".join(self._buf)))
            self._buf = []
        elif tag == "tr" and self._cells is not None:
            row = " | ".join(c for c in self._cells)
            self._cells = None
            if row.strip(" |"):
                self.blocks.append(Block("table", row))
        elif tag in self.BLOCK_TAGS:
            self._flush()

    def handle_data(self, data: str) -> None:
        if self._in_title:
            self.title += data
            return
        if self._skip or self._in_head:
            return
        self._buf.append(data if self._pre else re.sub(r"\s+", " ", data))

    def close(self) -> None:
        super().close()
        self._flush()


def html_to_blocks(html: str) -> tuple[list[Block], str]:
    parser = _HtmlText()
    parser.feed(html)
    parser.close()
    # Merge consecutive table rows into one table block.
    merged: list[Block] = []
    for block in parser.blocks:
        if block.kind == "table" and merged and merged[-1].kind == "table":
            merged[-1].text += "\n" + block.text
        else:
            merged.append(block)
    return merged, _clean(parser.title)


def _parse_html(data: bytes) -> ParsedDocument:
    blocks, title = html_to_blocks(decode_text(data))
    result = ParsedDocument(kind="html", parser="html", blocks=blocks)
    if title:
        result.metadata["title"] = title
    return result


# ---------------------------------------------------------------------------------------------- Markdown / text

_MD_HEADING = re.compile(r"^(#{1,6})\s+(.+?)\s*#*\s*$")


def markdown_to_blocks(text: str, page: int | None = None) -> tuple[list[Block], dict[str, Any]]:
    metadata: dict[str, Any] = {}
    lines = text.replace("\r\n", "\n").replace("\r", "\n").split("\n")
    if lines and lines[0].strip() == "---":
        for i in range(1, min(len(lines), 200)):
            if lines[i].strip() == "---":
                for meta_line in lines[1:i]:
                    key, sep, value = meta_line.partition(":")
                    if sep and key.strip():
                        metadata[key.strip()] = value.strip().strip("\"'")
                lines = lines[i + 1 :]
                break
    blocks: list[Block] = []
    buf: list[str] = []
    in_code = False
    code: list[str] = []

    def flush() -> None:
        content = _clean("\n".join(buf))
        buf.clear()
        if content:
            kind: BlockKind = "table" if content.startswith("|") else "text"
            blocks.append(Block(kind, content, page=page))

    for i, line in enumerate(lines):
        if line.strip().startswith("```"):
            if in_code:
                blocks.append(Block("code", "\n".join(code).strip("\n"), page=page))
                code = []
                in_code = False
            else:
                flush()
                in_code = True
            continue
        if in_code:
            code.append(line)
            continue
        match = _MD_HEADING.match(line)
        if match:
            flush()
            blocks.append(Block("heading", match.group(2).strip(), level=len(match.group(1)), page=page))
            continue
        nxt = lines[i + 1] if i + 1 < len(lines) else ""
        if line.strip() and re.fullmatch(r"=+\s*", nxt) and not buf:
            blocks.append(Block("heading", line.strip(), level=1, page=page))
            continue
        if re.fullmatch(r"=+\s*", line) and blocks and blocks[-1].kind == "heading":
            continue
        if not line.strip():
            flush()
            continue
        buf.append(line)
    if in_code and code:
        blocks.append(Block("code", "\n".join(code), page=page))
    flush()
    return blocks, metadata


def _parse_markdown(data: bytes) -> ParsedDocument:
    blocks, metadata = markdown_to_blocks(decode_text(data))
    return ParsedDocument(kind="markdown", parser="markdown", blocks=blocks, metadata=metadata)


def _parse_text(data: bytes, filename: str) -> ParsedDocument:
    text = decode_text(data)
    if filename.lower().endswith(".json"):
        try:
            text = json.dumps(json.loads(text), ensure_ascii=False, indent=2)
        except json.JSONDecodeError:
            pass
    blocks = [Block("text", p) for p in (_clean(p) for p in re.split(r"\n\s*\n", text)) if p]
    return ParsedDocument(kind="text", parser="text", blocks=blocks)


# ---------------------------------------------------------------------------------------------- e-mail


def _parse_eml(data: bytes, depth: int = 0) -> ParsedDocument:
    message = email.message_from_bytes(data, policy=email.policy.default)
    assert isinstance(message, EmailMessage)
    subject = str(message.get("subject", "") or "").strip()
    headers = {k: str(message.get(k, "") or "").strip() for k in ("from", "to", "cc", "date")}
    result = ParsedDocument(
        kind="eml",
        parser="email",
        metadata={"subject": subject, **{k: v for k, v in headers.items() if v}},
    )
    result.blocks.append(Block("heading", subject or "(sin asunto)", level=1))
    header_lines = [
        f"{label}: {headers[key]}"
        for key, label in (("from", "De"), ("to", "Para"), ("cc", "CC"), ("date", "Fecha"))
        if headers[key]
    ]
    if header_lines:
        result.blocks.append(Block("text", "\n".join(header_lines)))
    body = message.get_body(preferencelist=("plain", "html"))
    if body is not None:
        content = body.get_content()
        if isinstance(content, bytes):
            content = decode_text(content)
        if body.get_content_subtype() == "html":
            html_blocks, _ = html_to_blocks(str(content))
            result.blocks.extend(html_blocks)
        else:
            result.blocks.extend(Block("text", p) for p in (_clean(p) for p in re.split(r"\n\s*\n", str(content))) if p)
    attachments: list[str] = []
    for part in message.iter_attachments():
        name = part.get_filename() or "adjunto"
        attachments.append(name)
        if depth >= 2:
            continue
        payload = part.get_payload(decode=True)
        if not isinstance(payload, bytes) or not payload:
            continue
        try:
            kind = detect_kind(name, part.get_content_type(), payload)
            sub = _parse_sync(payload, name, part.get_content_type(), kind, depth + 1)
        except ParseError:
            continue
        result.blocks.append(Block("heading", f"Adjunto: {name}", level=2))
        for block in sub.blocks:
            result.blocks.append(Block(block.kind, block.text, level=block.level + 2 if block.level else 0))
        result.tables.extend(sub.tables)
    if attachments:
        result.metadata["attachments"] = attachments
    return result


# ---------------------------------------------------------------------------------------------- Docling


def docling_available() -> bool:
    if not get_settings().messer_kb_docling_enabled:
        return False
    try:
        import docling.document_converter  # type: ignore[import-not-found]  # noqa: F401
    except ImportError:
        return False
    return True


def _parse_with_docling(data: bytes, filename: str, kind: str) -> ParsedDocument | None:
    try:
        from docling.document_converter import DocumentConverter  # type: ignore[import-not-found]
    except ImportError:
        return None
    suffix = {"pdf": ".pdf", "docx": ".docx", "pptx": ".pptx", "html": ".html"}[kind]
    with tempfile.TemporaryDirectory(prefix="messer-docling-") as tmp:
        path = Path(tmp) / f"document{suffix}"
        path.write_bytes(data)
        converted: Any = DocumentConverter().convert(str(path))
    document: Any = converted.document
    result = ParsedDocument(kind=kind, parser="docling")
    pages = 0
    try:
        pages = int(document.num_pages())
    except Exception:
        pages = 0
    if pages > 0:
        for page_no in range(1, pages + 1):
            blocks, _ = markdown_to_blocks(document.export_to_markdown(page_no=page_no), page=page_no)
            result.blocks.extend(blocks)
        result.pages = pages
    else:
        blocks, _ = markdown_to_blocks(document.export_to_markdown())
        result.blocks.extend(blocks)
    return result


# ---------------------------------------------------------------------------------------------- entry points


def _parse_sync(data: bytes, filename: str, mime: str | None, kind: str, depth: int = 0) -> ParsedDocument:
    if kind in ("pdf", "docx", "pptx", "html") and depth == 0 and docling_available():
        try:
            parsed = _parse_with_docling(data, filename, kind)
            if parsed is not None and parsed.has_text:
                return parsed
        except Exception as exc:
            log.warning("parsing.docling_failed", filename=filename, error=str(exc)[:300])
    if kind == "pdf":
        return _parse_pdf(data)
    if kind == "docx":
        return _parse_docx(data)
    if kind == "pptx":
        return _parse_pptx(data)
    if kind == "xlsx":
        return _parse_xlsx(data, filename)
    if kind == "csv":
        return _parse_csv(data, filename)
    if kind == "html":
        return _parse_html(data)
    if kind == "markdown":
        return _parse_markdown(data)
    if kind == "eml":
        return _parse_eml(data, depth)
    return _parse_text(data, filename)


def parse_sync(data: bytes, filename: str, mime: str | None = None) -> ParsedDocument:
    kind = detect_kind(filename, mime, data)
    return _parse_sync(data, filename, mime, kind)


async def parse(data: bytes, filename: str, mime: str | None = None) -> ParsedDocument:
    """Parse a document off the event loop."""
    return await asyncio.to_thread(parse_sync, data, filename, mime)
